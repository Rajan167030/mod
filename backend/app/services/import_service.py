"""
Bulk data import from CSV or Excel into a scenario: aircraft, tasks, maintenance windows, sites,
crews, stores (weapons, countermeasures) and threats or airspace restrictions.

Two steps: preview (parse, map columns, validate every row, show what would change) and commit
(apply only when the file has no errors). Column names are matched loosely so files exported from
other systems can be used without editing headers.
"""
import csv
import io
import re
from datetime import datetime, timezone
from typing import Any, Callable, Dict, List, Tuple

from sqlalchemy.orm import Session

from app.core.errors import NotFoundError, ValidationError
from app.models.all_models import Resource, ResourceAvailability, Scenario, Task, TaskRequirement, User
from app.models.ops_models import Crew, OpsZone, StoreItem
from app.models.weather_models import Location
from app.services.audit_service import log_audit_event

MAX_ROWS = 2000
KINDS = ("aircraft", "tasks", "maintenance", "sites", "crews", "stores", "threats")

PRIORITIES = {"CRITICAL", "HIGH", "MEDIUM", "LOW"}
STATUSES = {"AVAILABLE", "MAINTENANCE", "GROUNDED", "UNAVAILABLE"}
SITE_TYPES = {"BASE", "WAYPOINT", "WEATHER_STATION", "POI", "LOGISTICS_HUB"}
CREW_STATUSES = {"AVAILABLE", "RESTING", "UNFIT"}
STORE_CATEGORIES = {"WEAPON", "COUNTERMEASURE", "MISSION_KIT"}
ZONE_KINDS = {"SAM": "THREAT", "AAA": "THREAT", "MANPADS": "THREAT", "RADAR": "THREAT", "RESTRICTED": "AIRSPACE", "NO_FLY": "AIRSPACE"}

# Canonical field -> accepted header spellings (compared after lower-casing and stripping punctuation)
ALIASES: Dict[str, Dict[str, List[str]]] = {
    "aircraft": {
        "name": ["name", "callsign", "aircraft", "tail", "tail number", "registration"],
        "type": ["type", "aircraft type", "resource type", "category"],
        "qualifications": ["qualifications", "capabilities", "roles", "quals"],
        "max_payload_kg": ["max payload kg", "payload kg", "payload", "max payload"],
        "status": ["status", "service status", "state"],
        "store_types": ["store types", "stores", "can carry", "carries", "weapons", "loadout"],
        "model": ["model", "type name", "aircraft model", "variant"],
        "service_ceiling_m": ["service ceiling m", "ceiling m", "service ceiling", "ceiling"],
        "range_km": ["range km", "range"],
        "cruise_speed_kmh": ["cruise speed kmh", "cruise speed", "cruise kmh"],
    },
    "tasks": {
        "name": ["name", "task", "mission", "task name", "title"],
        "priority": ["priority", "prio"],
        "duration_minutes": ["duration minutes", "duration", "minutes", "duration min"],
        "earliest_start": ["earliest start", "window open", "window opens", "start", "not before"],
        "latest_finish": ["latest finish", "deadline", "must finish by", "end", "not after", "finish by"],
        "qualification": ["qualification", "requires", "required qualification", "capability"],
        "site": ["site", "location", "destination", "site name"],
        "stores": ["stores", "store", "load", "weapons", "weapon load", "loadout"],
    },
    "maintenance": {
        "aircraft": ["aircraft", "name", "callsign", "tail"],
        "start": ["start", "from", "start time", "begins"],
        "end": ["end", "to", "until", "end time", "ends"],
    },
    "sites": {
        "name": ["name", "site", "location"],
        "latitude": ["latitude", "lat"],
        "longitude": ["longitude", "lon", "lng", "long"],
        "type": ["type", "site type", "kind"],
        "elevation_m": ["elevation m", "elevation", "altitude m", "alt m"],
    },
    "crews": {
        "name": ["name", "crew", "crew name", "callsign"],
        "aircraft_types": ["aircraft types", "aircraft type", "types", "current on", "flies", "type rating"],
        "status": ["status", "state", "fitness"],
        "max_duty_minutes": ["max duty minutes", "max duty", "duty limit", "duty limit min", "max duty min"],
        "min_rest_minutes": ["min rest minutes", "min rest", "rest", "rest min"],
        "duty_used_minutes": ["duty used minutes", "duty used", "flown", "hours flown min"],
        "duty_start": ["duty start", "shift start", "on duty"],
        "duty_end": ["duty end", "shift end", "off duty"],
    },
    "stores": {
        "store_type": ["store type", "type", "code", "store code"],
        "name": ["name", "store", "description", "item"],
        "category": ["category", "class"],
        "quantity": ["quantity", "qty", "count", "held", "stock"],
        "unit": ["unit", "units", "uom"],
    },
    "threats": {
        "name": ["name", "threat", "zone", "title"],
        "kind": ["kind", "type", "threat type", "zone type"],
        "effect": ["effect", "rule", "action"],
        "latitude": ["latitude", "lat", "center lat", "centre lat"],
        "longitude": ["longitude", "lon", "lng", "long", "center lon", "centre lon"],
        "radius_km": ["radius km", "radius", "range km", "range"],
        "polygon": ["polygon", "boundary", "points", "coordinates", "vertices"],
        "active_from": ["active from", "from", "start", "valid from"],
        "active_until": ["active until", "until", "end", "valid until", "to"],
        "exempt_capability": ["exempt capability", "exempt", "exemption", "allowed with"],
        "source": ["source", "intel source", "reported by"],
        "confidence": ["confidence", "conf", "confidence pct", "reliability"],
    },
}
REQUIRED = {
    "aircraft": ["name", "type"],
    "tasks": ["name", "priority", "duration_minutes", "earliest_start", "latest_finish"],
    "maintenance": ["aircraft", "start", "end"],
    "sites": ["name", "latitude", "longitude"],
    "crews": ["name", "aircraft_types"],
    "stores": ["store_type", "quantity"],
    # Geometry is either latitude + longitude + radius_km, or a polygon; checked per row
    "threats": ["name", "kind"],
}
TEMPLATES = {
    "aircraft": "name,type,qualifications,max_payload_kg,status,store_types\n"
                "C-17-Delta,HEAVY_AIRCRAFT,HEAVY_AIRLIFT;NIGHT_OPS;DASS,70000,AVAILABLE,FLARE_SET\n",
    "tasks": "name,priority,duration_minutes,earliest_start,latest_finish,qualification,site,stores\n"
             "WL-40 Bridge parts to Delta,HIGH,150,2026-10-17T06:00Z,2026-10-17T14:00Z,HEAVY_AIRLIFT,Delta Valley LZ,FLARE_SET\n",
    "crews": "name,aircraft_types,status,max_duty_minutes,min_rest_minutes,duty_used_minutes,duty_start,duty_end\n"
             "Globemaster crew,HEAVY_AIRCRAFT,AVAILABLE,480,60,0,,\n"
             "Night Mi-17 crew,TACTICAL_HELICOPTER,AVAILABLE,420,45,0,2026-10-17T12:00Z,2026-10-18T02:00Z\n",
    "stores": "store_type,name,category,quantity,unit\n"
              "FLARE_SET,Flare/chaff dispenser load,COUNTERMEASURE,8,sets\n"
              "ROCKET_POD,70 mm rocket pod,WEAPON,6,pods\n",
    "threats": "name,kind,effect,latitude,longitude,radius_km,active_from,active_until,exempt_capability,source,confidence\n"
               "SAM radar east of Echo,SAM,BLOCK,35.05,77.10,25,2026-10-17T08:00Z,2026-10-17T16:00Z,DASS,ELINT,70\n"
               "NOTAM R-220 gunnery,RESTRICTED,BLOCK,33.40,75.60,15,2026-10-17T10:00Z,2026-10-17T13:00Z,,NOTAM,100\n",
    "maintenance": "aircraft,start,end\nC-130-Bravo,2026-10-17T02:00Z,2026-10-17T06:00Z\n",
    "sites": "name,latitude,longitude,type,elevation_m\nGolf Supply Strip,33.90,76.10,BASE,3200\n",
}


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (h or "").lower()).strip()


def _read_rows(filename: str, data: bytes) -> List[Dict[str, Any]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise ValidationError(f"Could not read the Excel file: {e}")
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
    if name.endswith(".csv") or not name:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        if "\x00" in text:
            raise ValidationError("This does not look like a text CSV file.")
        rows = []
        try:
            for r in csv.DictReader(io.StringIO(text, newline="")):
                r.pop(None, None)  # cells beyond the header row have no column to go to
                if any(str(v or "").strip() for v in r.values()):
                    rows.append(r)
        except csv.Error as e:
            raise ValidationError(f"Could not read the CSV file: {e}")
        return rows
    raise ValidationError("Upload a .csv or .xlsx file.")


def _map_headers(kind: str, headers: List[str]) -> Tuple[Dict[str, str], List[str]]:
    """Canonical field -> original header, plus the headers that were not recognised."""
    mapping: Dict[str, str] = {}
    unused = []
    for h in headers:
        n = _norm(h)
        field = next((f for f, names in ALIASES[kind].items() if n in names or n == f.replace("_", " ")), None)
        if field and field not in mapping:
            mapping[field] = h
        elif h:
            unused.append(h)
    return mapping, unused


def _time(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    s = str(value).strip().replace(" ", "T", 1)
    if s.endswith(("Z", "z")):
        s = s[:-1] + "+00:00"
    dt = datetime.fromisoformat(s)
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _num(value: Any, kind=float) -> Any:
    if value in (None, ""):
        return None
    return kind(float(str(value).strip()))


def _list(value: Any) -> List[str]:
    return [q.strip().upper().replace(" ", "_") for q in re.split(r"[;|,/]", str(value or "")) if q.strip()]


# ---------------------------------------------------------------------------
# Row validators: return (clean values, errors)
# ---------------------------------------------------------------------------

def _check_aircraft(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["name"] = str(row.get("name") or "").strip()
    out["type"] = str(row.get("type") or "").strip().upper().replace(" ", "_")
    out["qualifications"] = _list(row.get("qualifications"))
    try:
        out["max_payload_kg"] = _num(row.get("max_payload_kg"), int)
    except ValueError:
        errs.append("max_payload_kg is not a number")
    status = str(row.get("status") or "AVAILABLE").strip().upper()
    if status not in STATUSES:
        errs.append(f"status must be one of {', '.join(sorted(STATUSES))}")
    out["status"] = status
    out["store_types"] = _list(row.get("store_types"))
    out["model"] = str(row.get("model") or "").strip()[:80] or None
    for f, hi in (("service_ceiling_m", 20000), ("range_km", 20000), ("cruise_speed_kmh", 1500)):
        try:
            out[f] = _num(row.get(f), int)
            if out[f] is not None and not 0 < out[f] <= hi:
                errs.append(f"{f} must be between 1 and {hi}")
        except ValueError:
            errs.append(f"{f} is not a number")
    out["_action"] = "update" if out["name"].lower() in ctx["aircraft"] else "create"
    return out, errs


def _check_task(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["name"] = str(row.get("name") or "").strip()
    prio = str(row.get("priority") or "").strip().upper()
    if prio not in PRIORITIES:
        errs.append(f"priority must be one of {', '.join(sorted(PRIORITIES))}")
    out["priority"] = prio
    try:
        out["duration_minutes"] = _num(row.get("duration_minutes"), int)
        if not out["duration_minutes"] or out["duration_minutes"] <= 0:
            errs.append("duration_minutes must be positive")
    except ValueError:
        errs.append("duration_minutes is not a number")
    for f in ("earliest_start", "latest_finish"):
        try:
            out[f] = _time(row.get(f))
        except (ValueError, TypeError):
            errs.append(f"{f} is not a date-time (use e.g. 2026-10-17T06:00Z)")
    if not errs and out["latest_finish"] <= out["earliest_start"]:
        errs.append("latest_finish must be after earliest_start")
    elif not errs and (out["latest_finish"] - out["earliest_start"]).total_seconds() / 60 < out["duration_minutes"]:
        errs.append("window is shorter than the duration")
    out["qualification"] = (_list(row.get("qualification")) or [None])[0]
    if out["qualification"] and out["qualification"] not in ctx["capabilities"]:
        ctx["warnings"].append(f"'{out['name']}': no aircraft holds {out['qualification']} yet")
    site = str(row.get("site") or "").strip()
    out["site"] = site or None
    if site and site.lower() not in ctx["sites"]:
        errs.append(f"site '{site}' is not in this scenario")
    out["stores"] = []
    for item in _list(row.get("stores")):
        kind, _, qty = item.partition(":")
        if qty and not qty.isdigit():
            errs.append(f"store '{item}' must look like TYPE or TYPE:quantity")
            continue
        out["stores"].append(f"{kind}:{qty}" if qty else kind)
        if kind not in ctx["store_types"]:
            ctx["warnings"].append(f"'{out['name']}': no {kind} held yet, so it cannot be flown until stores are imported")
    out["_action"] = "update" if out["name"].lower() in ctx["tasks"] else "create"
    return out, errs


def _int_in(row: Dict[str, Any], field: str, default: int, lo: int, hi: int, errs: List[str]) -> int:
    try:
        v = _num(row.get(field), int)
    except ValueError:
        errs.append(f"{field} is not a number")
        return default
    if v is None:
        return default
    if not lo <= v <= hi:
        errs.append(f"{field} must be between {lo} and {hi}")
    return v


def _opt_time(row: Dict[str, Any], field: str, errs: List[str]):
    if row.get(field) in (None, ""):
        return None
    try:
        return _time(row.get(field))
    except (ValueError, TypeError):
        errs.append(f"{field} is not a date-time (use e.g. 2026-10-17T06:00Z)")
        return None


def _check_crew(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["name"] = str(row.get("name") or "").strip()
    out["aircraft_types"] = _list(row.get("aircraft_types"))
    unknown = [t for t in out["aircraft_types"] if t not in ctx["aircraft_types"]]
    if unknown:
        ctx["warnings"].append(f"'{out['name']}': no aircraft of type {', '.join(unknown)} in this scenario")
    status = str(row.get("status") or "AVAILABLE").strip().upper()
    if status not in CREW_STATUSES:
        errs.append(f"status must be one of {', '.join(sorted(CREW_STATUSES))}")
    out["status"] = status
    out["max_duty_minutes"] = _int_in(row, "max_duty_minutes", 600, 30, 1440, errs)
    out["min_rest_minutes"] = _int_in(row, "min_rest_minutes", 60, 0, 1440, errs)
    out["duty_used_minutes"] = _int_in(row, "duty_used_minutes", 0, 0, 1440, errs)
    out["duty_start"] = _opt_time(row, "duty_start", errs)
    out["duty_end"] = _opt_time(row, "duty_end", errs)
    if out["duty_start"] and out["duty_end"] and out["duty_end"] <= out["duty_start"]:
        errs.append("duty_end must be after duty_start")
    out["_action"] = "update" if out["name"].lower() in ctx["crews"] else "create"
    return out, errs


def _check_store(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["store_type"] = str(row.get("store_type") or "").strip().upper().replace(" ", "_")
    out["name"] = str(row.get("name") or "").strip() or out["store_type"].replace("_", " ").title()
    cat = str(row.get("category") or "WEAPON").strip().upper().replace(" ", "_")
    if cat not in STORE_CATEGORIES:
        errs.append(f"category must be one of {', '.join(sorted(STORE_CATEGORIES))}")
    out["category"] = cat
    out["quantity"] = _int_in(row, "quantity", 0, 0, 100000, errs)
    out["unit"] = str(row.get("unit") or "each").strip()[:20]
    if out["store_type"] and not any(out["store_type"] in types for types in ctx["carriers"]):
        ctx["warnings"].append(f"{out['store_type']}: no aircraft can carry it yet (add it to an aircraft's store_types)")
    out["_action"] = "update" if out["store_type"] in ctx["store_types"] else "create"
    return out, errs


def _check_threat(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["name"] = str(row.get("name") or "").strip()
    kind = str(row.get("kind") or "").strip().upper().replace(" ", "_").replace("-", "_")
    kind = {"NOFLY": "NO_FLY", "NO_FLY_ZONE": "NO_FLY", "RESTRICTED_AIRSPACE": "RESTRICTED"}.get(kind, kind)
    if kind not in ZONE_KINDS:
        errs.append(f"kind must be one of {', '.join(ZONE_KINDS)}")
    out["kind"], out["category"] = kind, ZONE_KINDS.get(kind, "THREAT")
    effect = str(row.get("effect") or "BLOCK").strip().upper()
    effect = {"NO_ENTRY": "BLOCK", "NO ENTRY": "BLOCK", "AVOID_IF_POSSIBLE": "AVOID"}.get(effect, effect)
    if effect not in ("BLOCK", "AVOID"):
        errs.append("effect must be BLOCK or AVOID")
    out["effect"] = effect
    out["polygon"] = None
    out["latitude"] = out["longitude"] = out["radius_km"] = None
    if str(row.get("polygon") or "").strip():
        try:
            pts = [[float(a), float(b)] for a, b in (
                re.split(r"[\s,]+", p.strip()) for p in str(row["polygon"]).split(";") if p.strip()
            )]
            if len(pts) < 3 or any(not (-90 <= a <= 90 and -180 <= b <= 180) for a, b in pts):
                raise ValueError
            out["polygon"] = pts
        except ValueError:
            errs.append("polygon must be at least 3 points written as 'lat lon; lat lon; ...'")
    else:
        try:
            out["latitude"], out["longitude"] = _num(row.get("latitude")), _num(row.get("longitude"))
            if out["latitude"] is None or not -90 <= out["latitude"] <= 90:
                errs.append("latitude must be between -90 and 90 (or give a polygon)")
            if out["longitude"] is None or not -180 <= out["longitude"] <= 180:
                errs.append("longitude must be between -180 and 180 (or give a polygon)")
            out["radius_km"] = _num(row.get("radius_km"))
            if out["radius_km"] is None or not 0 < out["radius_km"] <= 500:
                errs.append("radius_km must be between 0 and 500 (or give a polygon)")
        except ValueError:
            errs.append("latitude, longitude and radius_km must be numbers")
    out["active_from"] = _opt_time(row, "active_from", errs)
    out["active_until"] = _opt_time(row, "active_until", errs)
    if out["active_from"] and out["active_until"] and out["active_until"] <= out["active_from"]:
        errs.append("active_until must be after active_from")
    exempt = str(row.get("exempt_capability") or "").strip().upper().replace(" ", "_")
    out["exempt_capability"] = exempt or None
    out["source"] = str(row.get("source") or "").strip()[:60] or None
    try:
        conf = _num(row.get("confidence"))
        conf = 0.8 if conf is None else conf / 100 if conf > 1 else conf  # accepts 70 or 0.7
        if not 0 <= conf <= 1:
            raise ValueError
        out["confidence"] = conf
    except ValueError:
        errs.append("confidence must be 0-1 or 0-100")
    out["_action"] = "update" if out["name"].lower() in ctx["zones"] else "create"
    return out, errs


def _check_maintenance(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    name = str(row.get("aircraft") or "").strip()
    out["resource_id"] = ctx["aircraft"].get(name.lower())
    out["aircraft"] = name
    if not out["resource_id"] and name.lower() not in ctx["pending_aircraft"]:
        errs.append(f"aircraft '{name}' is not in this scenario")
    for f in ("start", "end"):
        try:
            out[f] = _time(row.get(f))
        except (ValueError, TypeError):
            errs.append(f"{f} is not a date-time")
    if not errs and out["end"] <= out["start"]:
        errs.append("end must be after start")
    out["_action"] = "create"
    return out, errs


def _check_site(row: Dict[str, Any], ctx: Dict[str, Any]) -> Tuple[Dict[str, Any], List[str]]:
    errs, out = [], {}
    out["name"] = str(row.get("name") or "").strip()
    try:
        out["latitude"] = _num(row.get("latitude"))
        out["longitude"] = _num(row.get("longitude"))
        if out["latitude"] is None or not -90 <= out["latitude"] <= 90:
            errs.append("latitude must be between -90 and 90")
        if out["longitude"] is None or not -180 <= out["longitude"] <= 180:
            errs.append("longitude must be between -180 and 180")
    except ValueError:
        errs.append("latitude/longitude must be numbers")
    kind = str(row.get("type") or "POI").strip().upper().replace(" ", "_")
    if kind not in SITE_TYPES:
        errs.append(f"type must be one of {', '.join(sorted(SITE_TYPES))}")
    out["type"] = kind
    try:
        out["elevation_m"] = _num(row.get("elevation_m"), int)
    except ValueError:
        errs.append("elevation_m is not a number")
    out["_action"] = "update" if out["name"].lower() in ctx["sites"] else "create"
    return out, errs


CHECKERS: Dict[str, Callable] = {
    "aircraft": _check_aircraft,
    "tasks": _check_task,
    "maintenance": _check_maintenance,
    "sites": _check_site,
    "crews": _check_crew,
    "stores": _check_store,
    "threats": _check_threat,
}


def _context(db: Session, scenario_id: str) -> Dict[str, Any]:
    resources = db.query(Resource).filter(Resource.scenario_id == scenario_id).all()
    return {
        "aircraft": {r.name.lower(): r.id for r in resources},
        "pending_aircraft": set(),
        "capabilities": {q for r in resources for q in (r.metadata_json or {}).get("qualifications", [])},
        "tasks": {t.name.lower(): t.id for t in db.query(Task).filter(Task.scenario_id == scenario_id)},
        "sites": {loc.name.lower(): loc.id for loc in db.query(Location).filter(Location.scenario_id == scenario_id)},
        "aircraft_types": {r.resource_type for r in resources},
        "carriers": [set((r.metadata_json or {}).get("store_types", [])) for r in resources],
        "crews": {c.name.lower() for c in db.query(Crew).filter(Crew.scenario_id == scenario_id)},
        "store_types": {st.store_type for st in db.query(StoreItem).filter(StoreItem.scenario_id == scenario_id)},
        "zones": {z.name.lower() for z in db.query(OpsZone).filter(OpsZone.scenario_id == scenario_id)},
        "warnings": [],
    }


def analyse(db: Session, scenario_id: str, kind: str, filename: str, data: bytes) -> Dict[str, Any]:
    if kind not in KINDS:
        raise ValidationError(f"kind must be one of {', '.join(KINDS)}")
    if not db.query(Scenario).filter(Scenario.id == scenario_id).first():
        raise NotFoundError("Scenario", scenario_id)
    raw = _read_rows(filename, data)
    if len(raw) > MAX_ROWS:
        raise ValidationError(f"Files are limited to {MAX_ROWS} rows.")
    headers = list(raw[0].keys()) if raw else []
    mapping, unused = _map_headers(kind, headers)
    missing = [f for f in REQUIRED[kind] if f not in mapping]

    ctx = _context(db, scenario_id)
    rows, seen = [], set()
    for i, r in enumerate(raw, start=2):  # row 1 is the header
        values = {f: r.get(h) for f, h in mapping.items()}
        clean, errs = ({}, [f"missing column: {m}" for m in missing]) if missing else CHECKERS[kind](values, ctx)
        if not missing:
            for f in REQUIRED[kind]:
                if values.get(f) in (None, "") and not any(f in e for e in errs):
                    errs.append(f"{f} is empty")
            key = (clean.get("store_type") if kind == "stores" else clean.get("name") or clean.get("aircraft") or "").lower()
            if kind != "maintenance" and key in seen:
                errs.append("duplicate name in this file")
            seen.add(key)
            if kind == "aircraft" and key:
                ctx["pending_aircraft"].add(key)
        rows.append({"row": i, "values": {k: v for k, v in clean.items() if not k.startswith("_")},
                     "action": clean.get("_action"), "errors": errs})

    return {
        "kind": kind,
        "columns": mapping,
        "unused_columns": unused,
        "missing_columns": missing,
        "rows": rows,
        "summary": {
            "total": len(rows),
            "valid": sum(not r["errors"] for r in rows),
            "invalid": sum(bool(r["errors"]) for r in rows),
            "create": sum(r["action"] == "create" and not r["errors"] for r in rows),
            "update": sum(r["action"] == "update" and not r["errors"] for r in rows),
        },
        "warnings": ctx["warnings"][:20],
    }


def commit(db: Session, scenario_id: str, kind: str, filename: str, data: bytes, user: User) -> Dict[str, Any]:
    report = analyse(db, scenario_id, kind, filename, data)
    if report["missing_columns"] or report["summary"]["invalid"]:
        raise ValidationError("Fix the errors shown in the preview before importing.")

    created = updated = 0
    for r in report["rows"]:
        v = r["values"]
        if kind == "aircraft":
            res = db.query(Resource).filter(Resource.scenario_id == scenario_id, Resource.name.ilike(v["name"])).first()
            meta = {"qualifications": v["qualifications"], **({"max_payload_kg": v["max_payload_kg"]} if v.get("max_payload_kg") else {}),
                    **({"store_types": v["store_types"]} if v.get("store_types") else {}),
                    **{k: v[k] for k in ("model", "service_ceiling_m", "range_km", "cruise_speed_kmh") if v.get(k)}}
            if res:
                res.resource_type, res.status, res.metadata_json = v["type"], v["status"], {**(res.metadata_json or {}), **meta}
                updated += 1
            else:
                res = Resource(scenario_id=scenario_id, name=v["name"], resource_type=v["type"], status=v["status"], metadata_json=meta)
                db.add(res)
                db.flush()  # no availability rows means available throughout; import maintenance to block time
                created += 1
        elif kind == "tasks":
            task = db.query(Task).filter(Task.scenario_id == scenario_id, Task.name.ilike(v["name"])).first()
            site_id = (
                db.query(Location.id).filter(Location.scenario_id == scenario_id, Location.name.ilike(v["site"])).scalar()
                if v.get("site") else None
            )
            fields = dict(priority=v["priority"], duration_minutes=v["duration_minutes"], earliest_start=v["earliest_start"],
                          latest_finish=v["latest_finish"], location_id=site_id)
            if task:
                for k, val in fields.items():
                    setattr(task, k, val)
                db.query(TaskRequirement).filter(TaskRequirement.task_id == task.id).delete()
                updated += 1
            else:
                task = Task(scenario_id=scenario_id, name=v["name"], status="PENDING", description="Imported", **fields)
                db.add(task)
                created += 1
            db.flush()
            if v.get("qualification"):
                db.add(TaskRequirement(task_id=task.id, requirement_type="QUALIFICATION", requirement_value=v["qualification"]))
            for need in v.get("stores") or []:
                db.add(TaskRequirement(task_id=task.id, requirement_type="STORE", requirement_value=need))
        elif kind == "maintenance":
            rid = db.query(Resource.id).filter(Resource.scenario_id == scenario_id, Resource.name.ilike(v["aircraft"])).scalar()
            db.add(ResourceAvailability(resource_id=rid, start_at=v["start"], end_at=v["end"], available=False))
            created += 1
        elif kind == "crews":
            crew = db.query(Crew).filter(Crew.scenario_id == scenario_id, Crew.name.ilike(v["name"])).first()
            fields = {k: v[k] for k in ("aircraft_types", "status", "max_duty_minutes", "min_rest_minutes",
                                         "duty_used_minutes", "duty_start", "duty_end")}
            if crew:
                for k, val in fields.items():
                    setattr(crew, k, val)
                updated += 1
            else:
                db.add(Crew(scenario_id=scenario_id, name=v["name"], **fields))
                created += 1
        elif kind == "stores":
            item = db.query(StoreItem).filter(StoreItem.scenario_id == scenario_id, StoreItem.store_type == v["store_type"]).first()
            if item:
                item.name, item.category, item.quantity, item.unit = v["name"], v["category"], v["quantity"], v["unit"]
                updated += 1
            else:
                db.add(StoreItem(scenario_id=scenario_id, **v))
                created += 1
        elif kind == "threats":
            zone = db.query(OpsZone).filter(OpsZone.scenario_id == scenario_id, OpsZone.name.ilike(v["name"])).first()
            fields = dict(category=v["category"], kind=v["kind"], effect=v["effect"], center_lat=v["latitude"],
                          center_lon=v["longitude"], radius_km=v["radius_km"], polygon_json=v["polygon"],
                          active_from=v["active_from"], active_until=v["active_until"],
                          exempt_capability=v["exempt_capability"], source=v["source"], confidence=v["confidence"])
            if zone:
                for k, val in fields.items():
                    setattr(zone, k, val)
                updated += 1
            else:
                db.add(OpsZone(scenario_id=scenario_id, name=v["name"], **fields))
                created += 1
        else:
            loc = db.query(Location).filter(Location.scenario_id == scenario_id, Location.name.ilike(v["name"])).first()
            meta = {"elevation_m": v["elevation_m"]} if v.get("elevation_m") is not None else {}
            if loc:
                loc.latitude, loc.longitude, loc.location_type = v["latitude"], v["longitude"], v["type"]
                loc.metadata_json = {**(loc.metadata_json or {}), **meta}
                updated += 1
            else:
                db.add(Location(scenario_id=scenario_id, name=v["name"], latitude=v["latitude"], longitude=v["longitude"],
                                location_type=v["type"], metadata_json=meta))
                created += 1

    log_audit_event(
        db, action="DATA_IMPORTED", entity_type="SCENARIO", entity_id=scenario_id, actor=user,
        metadata={"kind": kind, "file": filename, "rows": report["summary"]["total"], "created": created, "updated": updated},
    )
    db.commit()
    return {"kind": kind, "created": created, "updated": updated, "rows": report["summary"]["total"]}
